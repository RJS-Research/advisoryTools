#!/usr/bin/env python3
"""Renderer A: python-pptx. Outline markdown to an editable .pptx with native charts.

    python3 render-python-pptx.py sample-outline.md -o out/a.pptx [--template brand.potx]
                                 [--brand /path/to/your/brand.json]

Charts are real PowerPoint chart objects, not images: a reader can click into the
data. That matters whenever a deck's argument turns on a number, because a figure
nobody can interrogate invites the suspicion that it was reverse engineered.

Branding comes from brand.json, never from constants in this file. Re-brand by
swapping that file, or point --brand at your own copy anywhere on disk. If
--template is given, its slide master supplies the look and brand.json only fills
what the template leaves open.
"""

import argparse
import importlib.util
import io
import json
import pathlib
import sys
import zipfile

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
except ImportError:
    # The one dependency that does not ship with macOS. A bare
    # ModuleNotFoundError here tells the user nothing they can act on, and the
    # Homebrew case below is the likeliest way the obvious install fails.
    sys.exit(
        "deckRender needs the python-pptx package, which does not ship with macOS.\n"
        "\n"
        "    python3 -m pip install --user python-pptx\n"
        "\n"
        "If that fails with 'error: externally-managed-environment', your python3 is\n"
        "Homebrew's and refuses to install into itself. Three ways out:\n"
        "\n"
        "  1. Use the python3 Apple ships, which has no such restriction:\n"
        "       /usr/bin/python3 -m pip install --user python-pptx\n"
        "       /usr/bin/python3 " + __file__ + " ...\n"
        "\n"
        "  2. Make a virtual environment, and render with ITS python:\n"
        "       python3 -m venv ~/.deckRender\n"
        "       ~/.deckRender/bin/pip install python-pptx\n"
        "       ~/.deckRender/bin/python " + __file__ + " ...\n"
        "\n"
        "  3. Install into your user site anyway. Simplest, but Homebrew may warn\n"
        "     about it when it upgrades python:\n"
        "       python3 -m pip install --user --break-system-packages python-pptx\n"
    )

HERE = pathlib.Path(__file__).resolve().parent

# The parser filename carries a hyphen, which is not importable, so load it by path
# rather than renaming the file.
_spec = importlib.util.spec_from_file_location("outline_parser", HERE / "outline-parser.py")
outline_parser = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(outline_parser)

CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}


def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, font_spec):
    run.font.name = font_spec["face"]
    run.font.size = Pt(font_spec["sizePt"])
    run.font.bold = font_spec.get("bold", False)
    run.font.italic = font_spec.get("italic", False)
    run.font.color.rgb = rgb(font_spec["color"])


def add_textbox(slide, text, font_spec, left, top, width, height, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        style_run(run, font_spec)
    return box


def add_chart(slide, spec, brand, left, top, width, height):
    data = CategoryChartData()
    data.categories = spec["categories"]
    for s in spec["series"]:
        data.add_series(s["name"], s["values"])

    ctype = CHART_TYPES.get(spec.get("type", "bar"))
    if ctype is None:
        raise ValueError(f"unsupported chart type {spec.get('type')!r}")

    gframe = slide.shapes.add_chart(ctype, left, top, width, height, data)
    chart = gframe.chart

    chart.has_title = bool(spec.get("title"))
    if chart.has_title:
        chart.chart_title.text_frame.text = spec["title"]
        for p in chart.chart_title.text_frame.paragraphs:
            for run in p.runs:
                style_run(run, brand["fonts"]["footnote"])

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False

    order = brand["chartSeriesOrder"]
    for i, series in enumerate(chart.plots[0].series):
        colour = rgb(order[i % len(order)])
        if ctype == XL_CHART_TYPE.LINE_MARKERS:
            series.format.line.color.rgb = colour
            series.format.line.width = Pt(2.25)
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = colour
    return chart


def open_template(template):
    """Open a .potx or .pptx template.

    python-pptx refuses a genuine .potx outright: pptx.api.Presentation checks the
    package content type and raises ValueError on
    presentationml.template.main+xml -- so "opens a real .potx" needs this step to
    be true at all, and a stand-in .pptx will pass without it.

    The two formats differ only in that content type, so a .potx is rewritten to
    the presentation type in memory. Nothing on disk is touched: brand.potx stays
    a real template that PowerPoint offers under New from Template.
    """
    path = pathlib.Path(template)
    if path.suffix.lower() != ".potx":
        return Presentation(str(path))

    template_ct = ("application/vnd.openxmlformats-officedocument"
                   ".presentationml.template.main+xml")
    presentation_ct = ("application/vnd.openxmlformats-officedocument"
                       ".presentationml.presentation.main+xml")

    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        payload = {n: zin.read(n) for n in names}
    ct_name = "[Content_Types].xml"
    ct = payload[ct_name].decode("utf-8")
    if template_ct not in ct:
        raise ValueError(f"{path.name}: not a template, no template content type")
    payload[ct_name] = ct.replace(template_ct, presentation_ct).encode("utf-8")
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            zout.writestr(name, payload[name])
    buf.seek(0)
    return Presentation(buf)


def render(outline_path, out_path, template=None, brand_path=None):
    # Default to the brand.json beside this script, but never require it: the copy
    # in an installed plugin lives in a versioned cache that `plugin update`
    # replaces, so a client's own palette has to be able to live somewhere else.
    brand = json.loads(pathlib.Path(brand_path or HERE / "brand.json").read_text())
    deck = outline_parser.parse(outline_path)
    L = brand["layout"]

    prs = open_template(template) if template else Presentation()
    if not template:
        prs.slide_width = Inches(L["widthIn"])
        prs.slide_height = Inches(L["heightIn"])

    blank = prs.slide_layouts[6]  # blank; we place every shape explicitly
    margin = Inches(L["marginIn"])
    content_w = prs.slide_width - 2 * margin

    # Title slide
    s = prs.slides.add_slide(blank)
    add_textbox(s, deck["title"], brand["fonts"]["heading"],
                margin, Inches(2.4), content_w, Inches(1.4))
    if deck["subtitle"]:
        add_textbox(s, deck["subtitle"], brand["fonts"]["body"],
                    margin, Inches(3.7), content_w, Inches(0.8))
    if deck["footer"]:
        add_textbox(s, deck["footer"], brand["fonts"]["footnote"],
                    margin, prs.slide_height - Inches(0.9), content_w, Inches(0.4))

    for sl in deck["slides"]:
        s = prs.slides.add_slide(blank)
        add_textbox(s, sl["title"], brand["fonts"]["slideTitle"],
                    margin, Inches(L["titleTopIn"]), content_w, Inches(0.9))

        bullets = [b for b in sl["blocks"] if b["kind"] == "bullet"]
        charts = [b for b in sl["blocks"] if b["kind"] == "chart"]

        # With a chart present the bullets take the left column, chart the right.
        body_w = content_w // 2 - Inches(0.15) if charts else content_w
        top = Inches(L["bodyTopIn"])

        if bullets:
            box = s.shapes.add_textbox(margin, top, body_w,
                                       prs.slide_height - top - Inches(0.8))
            tf = box.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = "•  " + b["text"]
                p.space_after = Pt(10)
                for run in p.runs:
                    style_run(run, brand["fonts"]["body"])

        for c in charts:
            left = margin + body_w + Inches(0.3) if bullets else margin
            width = content_w - body_w - Inches(0.3) if bullets else content_w
            add_chart(s, c["spec"], brand, left, top, width,
                      prs.slide_height - top - Inches(0.8))

        if sl["notes"]:
            s.notes_slide.notes_text_frame.text = sl["notes"]

    out_path = pathlib.Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return prs, out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outline")
    ap.add_argument("-o", "--out", default="out/python-pptx.pptx")
    ap.add_argument("--template", default=None, help="optional .potx/.pptx whose master supplies the look")
    ap.add_argument("--brand", default=None,
                    help="brand.json to use (default: the one beside this script)")
    a = ap.parse_args()

    prs, out = render(a.outline, a.out, a.template, a.brand)
    n_charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    n_notes = sum(1 for s in prs.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    print(f"python-pptx  -> {out}")
    print(f"  slides={len(prs.slides)}  native charts={n_charts}  slides with notes={n_notes}")
    print(f"  size={out.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
