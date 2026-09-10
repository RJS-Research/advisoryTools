#!/usr/bin/env python3
"""Build brand.potx from brand.json.

    python3 build-potx.py                  # writes brand.potx
    python3 build-potx.py -o other.potx

Why this exists. `render-python-pptx.py --template` opens a real .potx and
inherits its slide master, which is the one capability PptxGenJS does not have
(README, "The difference that actually decides it"). Until now the --template
path was proven only against a stand-in. This builds the real one.

What a template can and cannot do here. The renderer draws on `slide_layouts[6]`,
the blank layout, and places every shape explicitly. So the template does NOT
supply placeholder geometry. It supplies four things that still matter:

  1. Theme colour scheme, so a native chart, a hand-drawn shape, or anything a
     human adds later picks the brand palette out of PowerPoint's own colour
     picker rather than Office default blue.
  2. Theme font scheme, so +Headings and +Body resolve to Georgia and Arial.
  3. Master decoration, which shows through the blank layout: a baseline rule
     and the firm footer.
  4. Slide size, fixed at the brand's 16:9.

Method. python-pptx has no theme API, so the colour and font schemes are
surgically swapped inside ppt/theme/theme1.xml after save, leaving the format
scheme (fills, lines, effects) untouched. Authoring a whole theme from scratch
risks an invalid file for no gain. The package content type is then rewritten
from presentation to template, which is the only difference between .pptx and
.potx that PowerPoint actually reads.

Determinism. Same brand.json in, byte-identical .potx out, so the file can be
committed and diffed. Timestamps inside the zip are pinned for that reason.
"""

import argparse
import json
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent

# Fixed zip timestamp so the build is reproducible and the artefact diffs clean.
PINNED_DATE = (2026, 1, 1, 0, 0, 0)

PRESENTATION_CT = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
)
TEMPLATE_CT = (
    "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
)

# Theme slots mapped onto brand roles. dk1/lt1 are the text-background pair
# PowerPoint uses for automatic contrast; accent1..6 populate the chart cycle
# and the colour picker's top row, in brand.json's own chartSeriesOrder.
THEME_SLOTS = ["dk1", "lt1", "dk2", "lt2",
               "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
               "hlink", "folHlink"]


def load_brand(path):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def theme_colours(brand):
    """Return the twelve theme slots as hex strings, derived from brand.json."""
    pal = brand["palette"]
    roles = brand["roles"]
    series = brand["chartSeriesOrder"]

    # accent1..6 come off the top of the chart series order, so a chart rendered
    # against the theme and a chart rendered from brand.json agree.
    accents = (series + series)[:6]

    return {
        "dk1": roles["body"],            # 000000, automatic dark text
        "lt1": roles["background"],      # FFFFFF, automatic light background
        "dk2": roles["primaryText"],     # deep forest green, the brand's dark
        "lt2": pal["oliveBeige"],        # the light neutral
        "accent1": accents[0],
        "accent2": accents[1],
        "accent3": accents[2],
        "accent4": accents[3],
        "accent5": accents[4],
        "accent6": accents[5],
        "hlink": pal["steelBlue"],
        "folHlink": pal["warmWalnut"],
    }


def build_clr_scheme(colours, name="RJS Research"):
    """Build the <a:clrScheme> element text.

    dk1/lt1 are written as sysClr in the Office default; we write srgbClr for
    all twelve so the output does not depend on the host system palette.
    """
    parts = [f'<a:clrScheme name="{name}">']
    for slot in THEME_SLOTS:
        parts.append(f'<a:{slot}><a:srgbClr val="{colours[slot]}"/></a:{slot}>')
    parts.append("</a:clrScheme>")
    return "".join(parts)


def build_font_scheme(brand, name="RJS Research"):
    major = brand["fonts"]["heading"]["face"]
    minor = brand["fonts"]["body"]["face"]
    def block(tag, face):
        return (f'<a:{tag}>'
                f'<a:latin typeface="{face}"/>'
                f'<a:ea typeface=""/>'
                f'<a:cs typeface=""/>'
                f'</a:{tag}>')
    return (f'<a:fontScheme name="{name}">'
            + block("majorFont", major)
            + block("minorFont", minor)
            + "</a:fontScheme>")


def hexcolor(value):
    return RGBColor.from_string(value)


def decorate_master(prs, brand):
    """Add the master decoration that shows through the blank layout.

    Deliberately minimal. Two elements, both below the content area, so nothing
    the renderer places can collide with them. A heavier master would fight the
    renderer, which owns every content shape.

    python-pptx exposes add_shape and add_textbox on SlideShapes only; the
    MasterShapes and LayoutShapes collections are read-only in the public API.
    So the shapes are drawn on a scratch slide with the normal API, their XML is
    grafted into the master's shape tree, and the scratch slide is dropped along
    with its relationship so no orphan part survives into the template.
    """
    import copy

    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN

    master = prs.slide_master
    roles = brand["roles"]
    lay = brand["layout"]
    width = Inches(lay["widthIn"])
    height = Inches(lay["heightIn"])
    margin = Inches(lay["marginIn"])

    scratch = prs.slides.add_slide(prs.slide_layouts[6])

    # Baseline rule, sitting just above the footer.
    rule = scratch.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        margin, Emu(int(height - Inches(0.62))),
        Emu(int(width - 2 * margin)), Pt(1.0),
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = hexcolor(roles["accent"])
    rule.line.fill.background()
    rule.shadow.inherit = False
    rule.name = "brandRule"

    # Firm footer, left aligned under the rule.
    foot = scratch.shapes.add_textbox(
        margin, Emu(int(height - Inches(0.52))),
        Emu(int(width - 2 * margin)), Inches(0.32),
    )
    tf = foot.text_frame
    tf.word_wrap = False
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run()
    run.text = brand["_brand"]
    fnt = brand["fonts"]["footnote"]
    run.font.name = fnt["face"]
    run.font.size = Pt(fnt["sizePt"])
    run.font.bold = False
    run.font.color.rgb = hexcolor(fnt["color"])
    foot.name = "brandFooter"

    # Graft both shapes into the master, then discard the scratch slide.
    master_tree = master.shapes._spTree
    for shape in (rule, foot):
        master_tree.append(copy.deepcopy(shape._element))

    slide_ids = prs.slides._sldIdLst
    entry = slide_ids[-1]
    rel_id = entry.rId
    slide_ids.remove(entry)
    prs.part.drop_rel(rel_id)

    return master


def style_master_placeholders(prs, brand):
    """Point the master's title and body placeholders at the brand fonts.

    The renderer does not use these, but a human editing the deck in PowerPoint
    will, and a template whose placeholders are Office default while the drawn
    slides are Georgia is a template that produces inconsistent decks the first
    time someone adds a slide by hand.
    """
    master = prs.slide_master
    title_font = brand["fonts"]["slideTitle"]
    body_font = brand["fonts"]["body"]

    for ph in master.placeholders:
        idx_type = ph.placeholder_format.type
        tf = ph.text_frame
        target = title_font if "TITLE" in str(idx_type) else body_font
        for para in tf.paragraphs:
            para.font.name = target["face"]
            para.font.size = Pt(target["sizePt"])
            para.font.bold = bool(target.get("bold", False))
            para.font.color.rgb = hexcolor(target["color"])


def swap_theme(theme_xml, clr_scheme, font_scheme):
    """Replace the colour and font schemes inside an existing theme part."""
    out, n_clr = re.subn(r"<a:clrScheme.*?</a:clrScheme>", clr_scheme,
                         theme_xml, count=1, flags=re.S)
    if n_clr != 1:
        raise RuntimeError("theme1.xml: expected exactly one <a:clrScheme>")
    out, n_fnt = re.subn(r"<a:fontScheme.*?</a:fontScheme>", font_scheme,
                         out, count=1, flags=re.S)
    if n_fnt != 1:
        raise RuntimeError("theme1.xml: expected exactly one <a:fontScheme>")
    return out


def repackage(src_pptx, dest_potx, clr_scheme, font_scheme):
    """Rewrite the theme and the package content type, into a .potx."""
    with zipfile.ZipFile(src_pptx) as zin:
        names = zin.namelist()
        payload = {n: zin.read(n) for n in names}

    theme_parts = [n for n in names if re.fullmatch(r"ppt/theme/theme\d+\.xml", n)]
    if not theme_parts:
        raise RuntimeError("no ppt/theme/themeN.xml part found")
    for name in theme_parts:
        xml = payload[name].decode("utf-8")
        payload[name] = swap_theme(xml, clr_scheme, font_scheme).encode("utf-8")

    ct_name = "[Content_Types].xml"
    ct = payload[ct_name].decode("utf-8")
    if PRESENTATION_CT not in ct:
        raise RuntimeError("content types: presentation part type not found")
    payload[ct_name] = ct.replace(PRESENTATION_CT, TEMPLATE_CT).encode("utf-8")

    dest_potx.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(dest_potx, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in names:
            info = zipfile.ZipInfo(name, date_time=PINNED_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, payload[name])

    return theme_parts


def build(brand_path, out_path):
    brand = load_brand(brand_path)
    lay = brand["layout"]

    prs = Presentation()
    prs.slide_width = Inches(lay["widthIn"])
    prs.slide_height = Inches(lay["heightIn"])

    style_master_placeholders(prs, brand)
    decorate_master(prs, brand)

    colours = theme_colours(brand)
    clr_scheme = build_clr_scheme(colours)
    font_scheme = build_font_scheme(brand)

    tmpdir = Path(tempfile.mkdtemp())
    try:
        staged = tmpdir / "staged.pptx"
        prs.save(str(staged))
        themes = repackage(staged, out_path, clr_scheme, font_scheme)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    return brand, colours, themes


def main():
    ap = argparse.ArgumentParser(description="Build brand.potx from brand.json")
    ap.add_argument("--brand", default=str(HERE / "brand.json"))
    ap.add_argument("-o", "--out", default=str(HERE / "brand.potx"))
    a = ap.parse_args()

    out = Path(a.out)
    brand, colours, themes = build(Path(a.brand), out)

    print(f"brand.potx  -> {out}")
    print(f"  brand={brand['_brand']}  "
          f"size={brand['layout']['widthIn']}x{brand['layout']['heightIn']}in")
    print(f"  fonts: heading={brand['fonts']['heading']['face']}  "
          f"body={brand['fonts']['body']['face']}")
    print(f"  theme parts rewritten: {len(themes)}")
    print("  accents: " + " ".join(colours[f"accent{i}"] for i in range(1, 7)))
    print(f"  bytes={out.stat().st_size:,}")


if __name__ == "__main__":
    main()
